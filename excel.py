from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def exporter_excel(gestion, chemin):
    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats"

    headers = ["Matricule", "Nom", "Prénom", "Filière", "Moyenne", "Mention"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for etudiant in gestion.etudiants.values():
        moyenne, mention = gestion.resultat_etudiant(etudiant.matricule)
        ws.append([
            etudiant.matricule,
            etudiant.nom,
            etudiant.prenom,
            etudiant.filiere,
            round(moyenne, 2),
            mention
        ])

    # Deuxième feuille : détail des notes
    detail = wb.create_sheet("Notes")
    detail.append(["Matricule", "Étudiant", "Code", "Matière", "Coefficient", "Note"])
    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for n in gestion.notes:
        detail.append([
            n.etudiant.matricule,
            f"{n.etudiant.prenom} {n.etudiant.nom}",
            n.matiere.code,
            n.matiere.nom,
            n.matiere.coef,
            n.valeur
        ])

    for sheet in wb.worksheets:
        for column_cells in sheet.columns:
            max_len = max(len(str(c.value or "")) for c in column_cells)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_len + 2, 35)
        sheet.freeze_panes = "A2"

    wb.save(chemin)
